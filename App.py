import openpyxl

def lookup_values(file_path, sheet_name):
    try:
        # Load the workbook
        wb = openpyxl.load_workbook(file_path, data_only=True)
        if sheet_name not in wb.sheetnames:
            print(f"Sheet '{sheet_name}' not found.")
            return
        ws = wb[sheet_name]
    except Exception as e:
        print(f"Error loading workbook: {e}")
        return

    # Create a mapping from column 3 to column 2
    try:
        mapping = {}
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=3):
            col2_value = row[0].value
            col3_value = row[1].value
            if col3_value is not None:
                mapping[col3_value] = col2_value
    except Exception as e:
        print(f"Error creating mapping: {e}")
        return

    # Continuously prompt the user for input lists
    while True:
        print("Enter a list of numbers (comma-separated) or type 'done' to finish:")
        user_input = input()
        
        if user_input.lower() == 'done':
            print("\nExiting. Goodbye!")
            break

        print()  # One line break before "Results:"

        try:
            numbers = [int(num.strip()) for num in user_input.split(",")]
        except ValueError:
            print("Invalid input. Please enter a comma-separated list of numbers.")
            continue

        # Perform the lookup
        results = [mapping.get(num, "NA") for num in numbers]
        result_string = ", ".join(map(str, results))
        
        # Display the results
        print("Results:")
        print(result_string)
        print("\n~~~")  # Separator before the next prompt

# Example usage
file_path = "Uworld_Comlex2_Mapping.xlsx"
sheet_name = "comlex2_mapping (1)"

lookup_values(file_path, sheet_name)
